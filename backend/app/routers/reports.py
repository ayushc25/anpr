from collections import OrderedDict
from datetime import datetime
from io import BytesIO
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import text
from sqlalchemy.orm import Session
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font

from .. import models, schemas
from ..core.config import get_settings
from ..database import get_db
from ..deps import get_current_user

router = APIRouter(prefix="/reports", tags=["reports"])

VEHICLE_TYPES_FOR_SUMMARY = ["car", "motorbike", "bus", "truck", "bicycle"]


@router.get("/daily-summary", response_model=schemas.DailySummaryPage)
def daily_summary(
    db: Session = Depends(get_db),
    _: models.User = Depends(get_current_user),
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
):
    where_clauses = []
    params = {}
    if date_from:
        where_clauses.append("detected_at >= :date_from")
        params["date_from"] = date_from
    if date_to:
        where_clauses.append("detected_at <= :date_to")
        params["date_to"] = date_to
    where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

    # A "day" means a LOCAL day. Grouping the naive-UTC column directly put the
    # boundary at midnight UTC — 05:30 in IST — so every vehicle between local
    # midnight and 05:30 was reported on the previous day.
    params["display_tz"] = get_settings().locale.display_timezone
    local_day = "date_trunc('day', detected_at AT TIME ZONE 'UTC' AT TIME ZONE :display_tz)"

    total = db.execute(
        text(f"SELECT COUNT(DISTINCT {local_day}) FROM events {where_sql}"),
        params,
    ).scalar() or 0

    rows = db.execute(
        text(f"""
            SELECT {local_day} AS day,
                   COUNT(*) AS total,
                   COUNT(*) FILTER (WHERE direction = 'in_') AS entries,
                   COUNT(*) FILTER (WHERE direction = 'out_') AS exits,
                   COUNT(*) FILTER (WHERE status = 'registered') AS registered,
                   COUNT(*) FILTER (WHERE status = 'whitelist') AS whitelist,
                   COUNT(*) FILTER (WHERE status = 'blacklist') AS blacklist,
                   COUNT(*) FILTER (WHERE status = 'unknown') AS unknown
            FROM events
            {where_sql}
            GROUP BY day
            ORDER BY day DESC
            LIMIT :limit OFFSET :offset
        """),
        {**params, "limit": page_size, "offset": (page - 1) * page_size},
    ).all()

    items = [
        schemas.DailySummaryRow(
            date=r.day.strftime("%Y-%m-%d"),
            total=r.total, entries=r.entries, exits=r.exits,
            registered=r.registered, whitelist=r.whitelist,
            blacklist=r.blacklist, unknown=r.unknown,
        )
        for r in rows
    ]
    return schemas.DailySummaryPage(items=items, total=total, page=page, page_size=page_size)


def _autosize_columns(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)


def _write_detailed_sheet(ws, events):
    headers = [
        "Date & Time", "Plate Number", "Vehicle Type", "Vehicle Color", "Plate Color", "Camera", "Direction",
        "Vehicle Status", "Resident / Owner", "Flat Number", "Confidence", "OCR Confidence", "Image Reference",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for e in events:
        ws.append([
            e.detected_at.strftime("%Y-%m-%d %H:%M:%S"),
            e.plate_number,
            e.vehicle_type,
            (e.vehicle_color or "").title(),
            (e.plate_color or "").title(),
            e.camera.name if e.camera else "",
            e.direction.value if e.direction else "",
            e.status.value,
            e.vehicle.display_owner if e.vehicle else "",
            e.vehicle.display_flat if e.vehicle else "",
            round(e.detect_confidence or 0.0, 2),
            round(e.plate_confidence or 0.0, 2),
            e.vehicle_image_path or "",
        ])
    _autosize_columns(ws)


def _write_daily_summary_sheet(ws, events):
    """One row per calendar day: totals, entry/exit split, status
    breakdown, and a count per vehicle type - for a quick day-by-day
    traffic overview instead of scrolling through every raw event."""
    headers = (
        ["Date", "Total", "Entries", "Exits", "Registered", "Whitelist", "Blacklist", "Unknown"]
        + [t.title() for t in VEHICLE_TYPES_FOR_SUMMARY]
    )
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    days = OrderedDict()
    for e in events:
        day = e.detected_at.strftime("%Y-%m-%d")
        row = days.setdefault(day, {
            "total": 0, "entries": 0, "exits": 0,
            "registered": 0, "whitelist": 0, "blacklist": 0, "unknown": 0,
            **{t: 0 for t in VEHICLE_TYPES_FOR_SUMMARY},
        })
        row["total"] += 1
        if e.direction and e.direction.value == "in":
            row["entries"] += 1
        elif e.direction and e.direction.value == "out":
            row["exits"] += 1
        row[e.status.value] = row.get(e.status.value, 0) + 1
        if e.vehicle_type in row:
            row[e.vehicle_type] += 1

    for day in sorted(days.keys(), reverse=True):
        row = days[day]
        ws.append(
            [day, row["total"], row["entries"], row["exits"],
             row["registered"], row["whitelist"], row["blacklist"], row["unknown"]]
            + [row[t] for t in VEHICLE_TYPES_FOR_SUMMARY]
        )
    _autosize_columns(ws)


@router.get("/export")
def export_excel(
    db: Session = Depends(get_db),
    _: models.User = Depends(get_current_user),
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    status: Optional[str] = None,
    camera_id: Optional[int] = None,
    vehicle_type: Optional[str] = None,
    report_type: str = Query("detailed", pattern="^(detailed|daily)$"),
):
    query = db.query(models.Event)
    if date_from:
        query = query.filter(models.Event.detected_at >= date_from)
    if date_to:
        query = query.filter(models.Event.detected_at <= date_to)
    if status:
        query = query.filter(models.Event.status == status)
    if camera_id:
        query = query.filter(models.Event.camera_id == camera_id)
    if vehicle_type:
        query = query.filter(models.Event.vehicle_type == vehicle_type)
    events = query.order_by(models.Event.detected_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    if report_type == "daily":
        ws.title = "Daily Summary"
        _write_daily_summary_sheet(ws, events)
    else:
        ws.title = "ANPR Report"
        _write_detailed_sheet(ws, events)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    suffix = "daily_summary" if report_type == "daily" else "report"
    filename = f"anpr_{suffix}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
